from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageOps


ROOT = Path("/Users/don/Documents/career-ops")
SRC_PHOTO = ROOT / "mycv" / "职业照.jpg"
CROPPED = ROOT / "mycv" / "职业照_简历头像_圆角4x5.png"

DOCS = [
    (
        ROOT / "mycv" / "韦东波_简历.docx",
        ROOT / "mycv" / "韦东波_简历_带职业照.docx",
    ),
    (
        ROOT / "mycv" / "韦东波_机器人系统工程师_优化排版完整版.docx",
        ROOT / "mycv" / "韦东波_机器人系统工程师_带职业照.docx",
    ),
]


def crop_headshot() -> None:
    img = Image.open(SRC_PHOTO).convert("RGB")
    # Professional resume crop: face-centered 4:5 portrait with enough collar/shoulder context.
    crop_box = (265, 42, 805, 717)
    cropped = ImageOps.exif_transpose(img.crop(crop_box)).resize((540, 675), Image.LANCZOS)

    border = 10
    radius = 38
    canvas = Image.new("RGBA", (cropped.width + border * 2, cropped.height + border * 2), (255, 255, 255, 0))
    framed = Image.new("RGBA", canvas.size, (246, 248, 250, 255))
    photo_layer = Image.new("RGBA", canvas.size, (255, 255, 255, 0))
    photo_layer.paste(cropped, (border, border))

    mask = Image.new("L", canvas.size, 0)
    try:
        from PIL import ImageDraw

        draw = ImageDraw.Draw(mask)
        draw.rounded_rectangle([0, 0, canvas.width - 1, canvas.height - 1], radius=radius, fill=255)
        inner = Image.new("L", canvas.size, 0)
        inner_draw = ImageDraw.Draw(inner)
        inner_draw.rounded_rectangle(
            [border, border, canvas.width - border - 1, canvas.height - border - 1],
            radius=max(radius - border, 1),
            fill=255,
        )
    except Exception:
        inner = mask = Image.new("L", canvas.size, 255)

    canvas = Image.composite(framed, canvas, mask)
    canvas = Image.composite(photo_layer, canvas, inner)
    canvas.save(CROPPED, optimize=True)


def set_cell_width(cell, width_twips: int) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    for table in cell.tables:
        tbl = table._element
        tbl.getparent().remove(tbl)


def add_photo_to_cell(cell, width_inches=0.78) -> None:
    clear_cell(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.add_picture(str(CROPPED), width=Inches(width_inches))


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def insert_after(paragraph, element) -> None:
    paragraph._p.addnext(element)


def add_to_table_header(doc: Document) -> None:
    table = doc.tables[0]
    if len(table.rows[0].cells) < 2:
        return
    photo_cell = table.rows[0].cells[-1]
    set_cell_width(table.rows[0].cells[0], 8050)
    set_cell_width(photo_cell, 1310)
    add_photo_to_cell(photo_cell, 0.74)


def add_header_table_to_plain_doc(doc: Document) -> None:
    if len(doc.paragraphs) < 3:
        return

    name_p, title_p, contact_p = doc.paragraphs[0], doc.paragraphs[1], doc.paragraphs[2]
    header_texts = [name_p.text, title_p.text, contact_p.text]
    insert_anchor = contact_p

    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    left, right = table.rows[0].cells
    set_cell_width(left, 8050)
    set_cell_width(right, 1310)

    clear_cell(left)
    p = left.add_paragraph()
    run = p.add_run(header_texts[0])
    run.bold = True
    run.font.size = Pt(21)
    p.paragraph_format.space_after = Pt(2)

    p = left.add_paragraph()
    run = p.add_run(header_texts[1])
    run.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)

    p = left.add_paragraph()
    run = p.add_run(header_texts[2])
    run.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(0)

    add_photo_to_cell(right, 0.78)
    insert_after(insert_anchor, table._tbl)

    for paragraph in (name_p, title_p, contact_p):
        p = paragraph._element
        p.getparent().remove(p)


def remove_manual_page_breaks(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for br in paragraph._p.xpath(".//w:br[@w:type='page']"):
            br.getparent().remove(br)
        if not paragraph.text.strip() and not paragraph.runs:
            continue
        if not paragraph.text.strip() and "w:drawing" not in paragraph._p.xml:
            p = paragraph._element
            p.getparent().remove(p)


def compact_resume_spacing(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.52)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.58)
        section.right_margin = Inches(0.58)

    for paragraph in doc.paragraphs:
        fmt = paragraph.paragraph_format
        text = paragraph.text.strip()
        if not text:
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            continue
        if len(text) < 18 and not text.startswith(("•", "1.", "2.", "3.", "4.")):
            fmt.space_before = Pt(5)
            fmt.space_after = Pt(2)
            fmt.keep_with_next = True
        else:
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(2)
        if text.startswith("•"):
            fmt.line_spacing = 1.04
        else:
            fmt.line_spacing = 1.06


def main() -> None:
    crop_headshot()
    for src, out in DOCS:
        doc = Document(src)
        remove_manual_page_breaks(doc)
        if doc.tables:
            add_to_table_header(doc)
        else:
            add_header_table_to_plain_doc(doc)
        compact_resume_spacing(doc)
        doc.save(out)
        print(out)


if __name__ == "__main__":
    main()
