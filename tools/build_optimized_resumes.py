from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/don/Documents/career-ops")
PHOTO = ROOT / "mycv" / "职业照_简历头像_圆角4x5.png"
OUT_DIR = ROOT / "mycv"

BLUE = RGBColor(23, 79, 128)
GRAY = RGBColor(90, 90, 90)
LIGHT_GRAY = RGBColor(225, 230, 235)
BLACK = RGBColor(28, 28, 28)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "nil")


def clear_cell(cell) -> None:
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)


def set_run_font(run, size=9.4, bold=False, color=BLACK, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_para(paragraph, before=0, after=1.5, line=1.03):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_text(paragraph, text, size=9.4, bold=False, color=BLACK):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def add_header(doc, title, subtitle):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    remove_table_borders(table)
    left, right = table.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_width(left, 8050)
    set_cell_width(right, 1250)
    clear_cell(left)
    clear_cell(right)

    p = left.add_paragraph()
    set_para(p, after=1)
    add_text(p, "韦东波", size=22, bold=True)

    p = left.add_paragraph()
    set_para(p, after=1)
    add_text(p, title, size=10.4, color=BLACK)

    p = left.add_paragraph()
    set_para(p, after=0.5)
    add_text(p, subtitle, size=8.8, color=GRAY)

    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(PHOTO), width=Inches(0.74))


def add_section(doc, title):
    p = doc.add_paragraph()
    set_para(p, before=5.5, after=2)
    add_text(p, title, size=11.3, bold=True, color=BLUE)
    p.paragraph_format.keep_with_next = True

    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "B8C7D6")
    border.append(bottom)
    p_pr.append(border)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_para(p, after=1.4, line=1.02)
    p.paragraph_format.left_indent = Cm(0.36)
    p.paragraph_format.first_line_indent = Cm(-0.18)
    add_text(p, "• ", size=8.9)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, size=8.9, bold=True)
        add_text(p, text[len(bold_prefix) :], size=8.9)
    else:
        add_text(p, text, size=8.9)


def add_role(doc, company, role, date):
    p = doc.add_paragraph()
    set_para(p, before=2, after=1)
    add_text(p, company, size=9.7, bold=True)
    add_text(p, f"｜{role}", size=9.4, bold=True, color=BLUE)
    add_text(p, f"\t{date}", size=8.9, color=GRAY)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.4), WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.keep_with_next = True


def add_project(doc, title):
    p = doc.add_paragraph()
    set_para(p, before=2.2, after=1)
    add_text(p, title, size=9.4, bold=True)
    p.paragraph_format.keep_with_next = True


def add_skill_line(doc, label, text):
    p = doc.add_paragraph()
    set_para(p, after=1.2, line=1.02)
    add_text(p, label + "：", size=8.9, bold=True, color=BLUE)
    add_text(p, text, size=8.9)


def set_doc_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.43)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(9.2)


def add_footer(doc, label):
    section = doc.sections[0]
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=0)
    add_text(p, f"韦东波｜{label}", size=7.3, color=GRAY)


def build_robotics():
    doc = Document()
    set_doc_defaults(doc)
    add_header(
        doc,
        "机器人系统工程师｜EtherCAT / ROS2 / CANopen / 企业级 AI 系统架构",
        "深圳｜2 年工作经验｜181 7224 4940｜12132301@mail.sustech.edu.cn｜github.com/ZeroErrControl",
    )

    add_section(doc, "个人概述")
    p = doc.add_paragraph()
    set_para(p, after=1.5, line=1.02)
    add_text(
        p,
        "南方科技大学智能制造与机器人硕士，现任机器人关节模组企业机器人软件工程师 / 企业级 AI 负责人。工作覆盖关节通信 SDK、EtherCAT 稳定性测试、ROS2 / MoveIt 生态集成、企业级 RAG / Agent 系统建设，能把底层工程问题、客户现场问题和知识系统建设串成可交付闭环。",
        size=8.9,
    )

    add_section(doc, "核心能力")
    add_skill_line(doc, "机器人系统", "ROS2 / MoveIt / URDF / 关节控制 / EtherCAT / CANopen / CAN / CiA402 / PDO / SDO")
    add_skill_line(doc, "工业通信与稳定性", "SOEM / IGH EtherCAT / DC Sync / 实时 Linux / 多关节同步 / 掉 OP 问题排查")
    add_skill_line(doc, "AI 原生研发", "Codex / Claude Code / Cursor / 多 Agent 协作 / RAG / 评测体系 / 私有化模型部署")
    add_skill_line(doc, "项目推进", "跨部门协同 / 客户问题闭环 / 技术文档沉淀 / 开源生态 / NVIDIA 伙伴渠道")

    add_section(doc, "工作经历")
    add_role(doc, "深圳市零差云控科技有限公司", "机器人软件工程师 / 企业级 AI 负责人", "2024.07 - 至今")
    bullets = [
        "负责 eRob 关节通信 SDK 与软件生态建设，覆盖 CAN、CANopen、EtherCAT 等协议方向，推进 ROS2 / MoveIt / Demo / 技术文档沉淀。",
        "参与关节 EtherCAT 通讯层运行稳定性测试，围绕掉 OP、首次启动失败、多关节同步等问题开展复现、定位与方案推进。",
        "从 0 到 1 主导企业级 AI 系统建设，覆盖机房、模型部署、全栈架构、数据清洗、知识库、RAG 与私有评测体系。",
        "推动售前售后重复产品咨询减少约 50%；自研 RAG 在内部测试集 Top-10 召回率达到 98.0%，高于行业平均水平约 91%。",
        "负责 NVIDIA Inception 申报、路演与生态合作推进，推动公司在 340 家参评企业中进入 Top 10 并获荣耀企业认证。",
    ]
    for b in bullets:
        add_bullet(doc, b)

    add_section(doc, "代表项目")
    add_project(doc, "1. eRob 关节模组通信 SDK 与开源生态建设｜项目负责人 / 核心开发")
    for b in [
        "面向关节模组客户建设统一 SDK 与示例工程，覆盖 EtherCAT、CANopen、CAN 等主要通信方向。",
        "沉淀设备初始化、CiA402 状态机、控制模式切换、PDO / SDO 通信、故障处理等可复用工程能力。",
        "完善 GitHub 开源仓库、技术文档与 Demo 工程，仓库周下载量超过 50 次，降低重复技术支持成本。",
    ]:
        add_bullet(doc, b)
    add_project(doc, "2. EtherCAT 通信稳定性与多关节压力测试｜问题分析 / 闭环推进")
    for b in [
        "围绕 SOEM / IGH 主站、DC 同步、SM 同步、实时 Linux、Intel I350 网卡和多从站调度排查掉 OP / 主站掉线问题。",
        "将客户现场异常转化为可复现测试链路，协同研发、技术支持和客户形成可执行解决方案。",
    ]:
        add_bullet(doc, b)
    add_project(doc, "3. ZeroErr GPT 企业级 AI 售前售后系统｜系统负责人 / 架构与核心研发")
    for b in [
        "建设企业专有知识服务底座，覆盖产品手册、通信协议、FAQ、售后经验、表格和图片资料。",
        "设计文档解析、知识切分、混合检索、重排、Evidence API、自动化评测与失败样本回流机制。",
        "支持 OpenAI-compatible API、OpenWebUI、Dify、LangChain、Codex 等生态接入，形成售前售后统一入口。",
    ]:
        add_bullet(doc, b)
    add_project(doc, "4. 机器人仿真与 NVIDIA / Isaac 生态探索｜方案验证 / 技术预研")
    for b in [
        "参与 URDF / STL 模型整理、Isaac Sim / Isaac Lab / ROS2 / MoveIt 集成探索，支持公司机器人关节进入 NVIDIA 生态。",
        "配合 AI 展区机械臂项目调试、实时内核环境安装与交付支持，为具身智能方向预研提供基础。",
    ]:
        add_bullet(doc, b)

    add_section(doc, "教育背景与成果")
    add_bullet(doc, "南方科技大学｜智能制造与机器人 · 硕士｜2021 - 2024.07；研究方向：长时储能与机器学习。")
    add_bullet(doc, "广西科技大学｜车辆工程 · 本科｜2016 - 2020。")
    add_bullet(doc, "共同第一作者在 Energy & Environmental Science 等顶级期刊发表论文 2 篇；已授权国家发明专利 2 项、实用新型专利 2 项。")
    add_bullet(doc, "全国大学生智能车竞赛全国二等奖（组内排名第一）；全国大学生励志奖学金。")

    add_footer(doc, "机器人系统工程师定向简历")
    out = OUT_DIR / "韦东波_机器人系统工程师_模板优化版.docx"
    doc.save(out)
    return out


def build_ai_robotics():
    doc = Document()
    set_doc_defaults(doc)
    add_header(
        doc,
        "AI应用 / 机器人生态工程师｜RAG / Agent / SDK",
        "深圳｜2 年工作经验｜181 7224 4940｜12132301@mail.sustech.edu.cn｜github.com/ZeroErrControl",
    )

    add_section(doc, "个人概述")
    p = doc.add_paragraph()
    set_para(p, after=1.5, line=1.02)
    add_text(
        p,
        "机器人关节模组企业 AI 系统负责人 / 机器人软件工程师，具备“机器人底层系统 + 企业级 AI 应用 + 开发者生态”的复合经验。能从 0 到 1 搭建 RAG / Agent / 知识库 / SDK / 文档体系，把客户问题、产品资料和研发经验转化为可检索、可评测、可复用的工程基础设施。",
        size=8.9,
    )

    add_section(doc, "能力关键词")
    add_skill_line(doc, "AI 应用与数据基建", "RAG / Agent / 文档解析 / 知识切分 / 混合检索 / 重排 / Evidence API / 自动化评测")
    add_skill_line(doc, "工程落地", "Python / FastAPI / CLI / OpenAI-compatible API / OpenWebUI / Dify / LangChain / 私有化模型部署")
    add_skill_line(doc, "机器人生态", "EtherCAT / CANopen / CAN / ROS2 / MoveIt / URDF / SDK / 技术文档 / 开源案例")
    add_skill_line(doc, "AI 工具研发能力", "Codex / Claude Code / Cursor / Manus / 多 Agent 协作 / 复杂问题拆解 / 代码生成与审查")

    add_section(doc, "工作经历")
    add_role(doc, "深圳市零差云控科技有限公司", "AI 系统负责人 / 机器人软件工程师", "2024.07 - 至今")
    for b in [
        "主导 ZeroErr GPT / TechChat 企业级售后知识问答系统，从基础设施、模型部署、数据清洗到 RAG 核心链路全流程落地。",
        "将产品手册、通信协议、FAQ、售后经验、表格和图片资料治理为可检索、可引用、可评测的企业知识服务。",
        "建立私有测试集和自动化评测体系，自研 RAG Top-10 召回率达到 98.0%，并通过失败样本回流持续优化。",
        "负责 eRob 关节通信 SDK 和开源生态建设，覆盖 EtherCAT / CANopen / CAN / ROS2 / MoveIt 等方向。",
        "推动 NVIDIA Inception 生态合作，公司从 340 家参评企业进入 Top 10，获得荣耀企业认证并提升产业客户曝光。",
    ]:
        add_bullet(doc, b)

    add_section(doc, "代表项目")
    add_project(doc, "1. ZeroErr GPT / TechChat 企业级售后知识问答系统｜AI 系统负责人")
    for b in [
        "设计“知识治理 + RAG + Evidence Tool + 自动化评测 + 业务闭环”的系统架构，支撑客服、研发和客户自助查询。",
        "构建文档解析、标题增强、章节结构、关联问题增强、多路检索、重排和证据返回链路。",
        "封装 OpenAI-compatible API 和 CLI，支持 OpenWebUI、Dify、LangChain、Codex、OpenClaw 等生态接入。",
        "将售前报价/选型/交期、售后问答、研发资料沉淀到统一入口，减少消息渠道分散和人工经验依赖。",
    ]:
        add_bullet(doc, b)
    add_project(doc, "2. eRob 机器人关节软件生态与通信 SDK｜项目负责人 / 开源维护")
    for b in [
        "建设面向客户和开发者的软件生态，覆盖 EtherCAT、CANopen、CAN、ROS2、MoveIt、URDF、Isaac Sim 等接口与示例。",
        "沉淀 CiA402 状态机、PDO / SDO、CSP / CST / PP / PT 控制模式、多关节同步和实时 Linux 调试经验。",
        "把客户现场问题转化为文档、示例、FAQ、技术支持话术和知识库数据，形成可复用的开发者支持体系。",
    ]:
        add_bullet(doc, b)
    add_project(doc, "3. 企业数字化与 AI 原生组织建设｜方案设计 / 推动落地")
    for b in [
        "围绕数据治理、知识库、数据中心、AI Agent 协作和业务流程线上化进行方案设计。",
        "探索部门、岗位、员工、Agent、人类协作主体之间的系统架构，推动 AI 系统从单点问答工具走向组织级知识基础设施。",
    ]:
        add_bullet(doc, b)
    add_project(doc, "4. 机器人仿真与 NVIDIA / Isaac 生态探索｜生态集成 / 展示支持")
    for b in [
        "参与 URDF / STL 模型整理、Isaac Sim / Isaac Lab / ROS2 / MoveIt 集成探索，支持机器人关节在 NVIDIA 平台展示。",
        "关注 VLA、世界模型、强化学习和具身智能方向，为 AI 机器人应用和数据基建方向储备工程场景。",
    ]:
        add_bullet(doc, b)

    add_section(doc, "教育背景与成果")
    add_bullet(doc, "南方科技大学｜智能制造与机器人 · 硕士｜2021 - 2024.07；研究方向：长时储能与机器学习。")
    add_bullet(doc, "广西科技大学｜车辆工程 · 本科｜2016 - 2020。")
    add_bullet(doc, "共同第一作者发表 EES 等顶刊论文 2 篇；授权国家发明专利 2 项、实用新型专利 2 项。")
    add_bullet(doc, "NVIDIA Inception 2025 荣耀企业 Top 10 / 340；全国大学生智能车竞赛全国二等奖。")

    add_footer(doc, "AI + 机器人复合方向简历")
    out = OUT_DIR / "韦东波_AI机器人复合方向_模板优化版.docx"
    doc.save(out)
    return out


def main():
    for out in (build_robotics(), build_ai_robotics()):
        print(out)


if __name__ == "__main__":
    main()
